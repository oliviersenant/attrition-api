"""Construit le support de soutenance (.pptx) de la mission Déploiement (Futurisys).

Agrège les Parties 2 « Soutenance » des synthèses. Exécution :
    uv run --with python-pptx --with matplotlib python presentation/build_deck.py
"""

import math
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ICI = Path(__file__).parent

# --- Palette (bleu nuit exécutif) -------------------------------------------
NAVY = RGBColor(0x1B, 0x2A, 0x4A)
NAVY2 = RGBColor(0x24, 0x3A, 0x63)
LIGHT = RGBColor(0xF5, 0xF6, 0xF8)
BLUE = RGBColor(0xCA, 0xDC, 0xFC)
ACCENT = RGBColor(0xE8, 0x6A, 0x33)
GREEN = RGBColor(0x2E, 0x8B, 0x57)
INK = RGBColor(0x1C, 0x24, 0x33)
GREY = RGBColor(0x5C, 0x66, 0x75)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

HEAD = "Cambria"
BODY = "Calibri"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
W, H = prs.slide_width, prs.slide_height


def slide(bg=LIGHT):
    s = prs.slides.add_slide(BLANK)
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, H)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def box(s, text, left, top, width, height, size, *, bold=False, color=INK,
        align=PP_ALIGN.LEFT, font=BODY, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(0)
    tf.margin_top = tf.margin_bottom = Pt(0)
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        r = p.add_run()
        r.text = line
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.name = font
        r.font.color.rgb = color
    return tb


def rect(s, left, top, width, height, color, line=None):
    sp = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top),
                            Inches(width), Inches(height))
    sp.fill.solid()
    sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(1)
    sp.shadow.inherit = False
    sp.adjustments[0] = 0.06
    return sp


def circle_num(s, n, left, top, d=0.62, fill=ACCENT):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(top), Inches(d), Inches(d))
    c.fill.solid()
    c.fill.fore_color.rgb = fill
    c.line.fill.background()
    c.shadow.inherit = False
    tf = c.text_frame
    tf.margin_top = tf.margin_bottom = Pt(0)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.name = HEAD
    r.font.color.rgb = WHITE
    return c


def header(s, kicker, title):
    circle_num(s, str(kicker).lstrip("0") or "0", 0.8, 0.62)
    box(s, title, 1.65, 0.55, 11.0, 0.9, 30, bold=True, color=NAVY, font=HEAD,
        anchor=MSO_ANCHOR.MIDDLE)


def bullets(s, items, left, top, width, size=15, gap=0.14, color=INK, lead=ACCENT):
    y = top
    sub_fs = size - 2.5
    for head_txt, sub in items:
        dot = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(left), Inches(y + 0.07),
                                 Inches(0.12), Inches(0.12))
        dot.fill.solid()
        dot.fill.fore_color.rgb = lead
        dot.line.fill.background()
        dot.shadow.inherit = False
        box(s, head_txt, left + 0.32, y, width - 0.32, 0.35, size, bold=True, color=color)
        h = 0.32
        if sub:
            cpl = max(12, int((width - 0.32) * 138 / sub_fs))  # chars par ligne (approx.)
            lines = max(1, math.ceil(len(sub) / cpl))
            line_h = sub_fs * 1.3 / 72
            box(s, sub, left + 0.32, y + h, width - 0.32, lines * line_h + 0.1, sub_fs,
                color=GREY, spacing=1.05)
            h += lines * line_h + 0.08
        y += h + gap
    return y


def card(s, left, top, width, height, title, value, note, accent=ACCENT):
    rect(s, left, top, width, height, WHITE, line=RGBColor(0xE1, 0xE5, 0xEC))
    rect(s, left, top, 0.09, height, accent)
    box(s, value, left + 0.28, top + 0.16, width - 0.4, 0.6, 30, bold=True, color=NAVY, font=HEAD)
    box(s, title, left + 0.28, top + 0.78, width - 0.4, 0.35, 12.5, bold=True, color=INK)
    box(s, note, left + 0.28, top + 1.12, width - 0.4, height - 1.15, 10.5, color=GREY, spacing=1.0)


# ============================================================ 1. TITRE
s = slide(NAVY)
rect(s, 0, 0, W.inches, 0.28, ACCENT)
box(s, "FUTURISYS · MISSION DE DÉPLOIEMENT MLOps", 0.9, 1.7, 11.5, 0.5, 15, bold=True,
    color=BLUE, font=BODY)
box(s, "Attrition API", 0.9, 2.35, 11.5, 1.2, 54, bold=True, color=WHITE, font=HEAD)
box(s, "Déployer un modèle de Machine Learning en production", 0.9, 3.7, 11.5, 0.7, 24,
    color=BLUE, font=HEAD)
box(s, "Du notebook d'analyse à une API testée, tracée et déployée en continu",
    0.9, 4.45, 11.5, 0.5, 15, color=RGBColor(0xAE, 0xC2, 0xE8))
rect(s, 0.9, 5.4, 4.6, 0.02, NAVY2)
box(s, "Olivier Senant  ·  Parcours AI/ML Engineer  ·  Soutenance",
    0.9, 5.6, 11.5, 0.4, 13, color=BLUE)
box(s, "attrition-api.onrender.com/docs", 0.9, 6.05, 11.5, 0.4, 13, bold=True, color=ACCENT)

# ============================================================ 2. CONTEXTE
s = slide()
header(s, "01", "Contexte : rendre un modèle opérationnel")
box(s, "Futurisys veut exploiter en production le modèle de prédiction d'attrition "
       "issu de la mission d'analyse (TechNova). On ne refait pas de data science : "
       "on transforme un modèle de notebook en service fiable et maintenable.",
    0.9, 1.7, 11.5, 0.9, 16, color=INK, spacing=1.15)
bullets(s, [
    ("Un dépôt Git structuré, versionné, collaboratif", "branches, tags, PR, README, conventions"),
    ("Une API performante exposant le modèle", "FastAPI + Pydantic, documentée (Swagger)"),
    ("Fiabilité garantie par les tests", "Pytest + rapport de couverture"),
    ("Traçabilité de tous les échanges", "base PostgreSQL"),
    ("Livraison automatisée", "pipeline CI/CD, gestion des secrets et environnements"),
], 0.95, 2.75, 7.2, size=15, gap=0.12)
# carte "mission"
rect(s, 8.7, 2.7, 3.8, 3.4, NAVY)
box(s, "LE DÉFI", 9.0, 2.95, 3.2, 0.4, 13, bold=True, color=ACCENT)
box(s, "Un modèle qui ne vit que dans un notebook n'a aucune valeur opérationnelle.",
    9.0, 3.4, 3.2, 1.2, 15, color=WHITE, spacing=1.15)
box(s, "Objectif : le rendre interrogeable par n'importe quel système, 24/7, "
       "avec les garanties d'un logiciel de production.",
    9.0, 4.55, 3.2, 1.4, 13, color=BLUE, spacing=1.15)

# ============================================================ 3. ARCHITECTURE
s = slide(NAVY)
box(s, "Architecture de bout en bout", 0.8, 0.6, 11.5, 0.8, 30, bold=True, color=WHITE,
    font=HEAD)
box(s, "Chaîne CI/CD complète : du push au service en ligne, sans intervention manuelle",
    0.8, 1.4, 11.5, 0.5, 15, color=BLUE)

flow = [
    ("Développeur", "branche → PR", BLUE),
    ("GitHub Actions", "lint · 54 tests · couverture", BLUE),
    ("Tag de version", "déclenche le déploiement", ACCENT),
    ("Render (Docker)", "build → ré-entraîne → sert", GREEN),
    ("API live", "FastAPI + Swagger", BLUE),
]
x = 0.8
wid = 2.15
for i, (t, sub, col) in enumerate(flow):
    rect(s, x, 2.5, wid, 1.5, NAVY2, line=col)
    box(s, t, x + 0.12, 2.72, wid - 0.24, 0.6, 15, bold=True, color=WHITE,
        align=PP_ALIGN.CENTER, font=HEAD)
    box(s, sub, x + 0.12, 3.35, wid - 0.24, 0.55, 11, color=col, align=PP_ALIGN.CENTER)
    if i < len(flow) - 1:
        box(s, "▶", x + wid - 0.02, 3.0, 0.5, 0.5, 18, color=ACCENT, align=PP_ALIGN.CENTER)
    x += wid + 0.28

# base
rect(s, 4.2, 4.9, 4.9, 1.15, NAVY2, line=ACCENT)
box(s, "PostgreSQL managé (Neon)", 4.4, 5.08, 4.5, 0.4, 15, bold=True, color=WHITE,
    align=PP_ALIGN.CENTER, font=HEAD)
box(s, "chaque prédiction est tracée : input · output · seuil · version · horodatage",
    4.4, 5.5, 4.5, 0.5, 11, color=BLUE, align=PP_ALIGN.CENTER)
box(s, "▲  DATABASE_URL (secret)", 6.0, 4.55, 3.0, 0.35, 12, color=ACCENT, bold=True)
box(s, "Secrets injectés par l'hébergeur — jamais dans le dépôt.",
    0.8, 6.4, 11.5, 0.4, 13, color=BLUE, align=PP_ALIGN.CENTER)

# ============================================================ 4. GIT
s = slide()
header(s, "02", "Dépôt Git : un historique qui raconte la mission")
bullets(s, [
    ("main ne porte que du code déployable", "une fonctionnalité = une branche, fusionnée par Pull Request après tests"),
    ("Conventions explicites (CONTRIBUTING.md)", "préfixes feat/ fix/ ci/ docs/ test/ ; commits descriptifs"),
    ("Versionnage sémantique par tags", "v0.1.0 → v1.0.0 ; un tag = une version livrable = un déploiement"),
    ("Structure claire, rien à la racine", "app/ · ml/ · scripts/ · tests/ · docs/"),
], 0.95, 1.85, 7.0, size=15, gap=0.16)
for i, (v, t) in enumerate([("7", "tags de version"), ("9", "Pull Requests"),
                            ("6", "étapes livrées")]):
    card(s, 8.5, 1.9 + i * 1.55, 4.0, 1.4, t, v, "", accent=NAVY)

# ============================================================ 5. CI/CD
s = slide()
header(s, "03", "CI/CD : la qualité devient une loi, pas une intention")
bullets(s, [
    ("Intégration continue (GitHub Actions)", "à chaque push : ruff + 54 tests + couverture, en < 10 min"),
    ("Protection de branche", "merge dans main impossible tant que le check n'est pas vert"),
    ("Déploiement continu par tag", "git push v1.0.0 → tests → deploy hook Render (moindre privilège)"),
    ("Environnements & secrets", "environnement production GitHub ; secrets jamais versionnés"),
], 0.95, 1.85, 7.1, size=15, gap=0.18)
rect(s, 8.5, 1.95, 4.0, 4.3, NAVY)
box(s, "2 fichiers YAML", 8.8, 2.2, 3.5, 0.4, 16, bold=True, color=ACCENT, font=HEAD)
box(s, "ci.yml", 8.8, 2.75, 3.5, 0.35, 14, bold=True, color=WHITE, font="Courier New")
box(s, "tests, couverture, déclenchement du déploiement",
    8.8, 3.1, 3.5, 0.6, 12, color=BLUE, spacing=1.1)
box(s, "render.yaml", 8.8, 3.85, 3.5, 0.35, 14, bold=True, color=WHITE, font="Courier New")
box(s, "Infrastructure-as-Code : le service décrit en versionné",
    8.8, 4.2, 3.5, 0.7, 12, color=BLUE, spacing=1.1)
rect(s, 8.8, 5.05, 3.4, 0.02, NAVY2)
box(s, "Split standard :\nGitHub = qualité · Render = livraison",
    8.8, 5.2, 3.5, 0.9, 12.5, color=WHITE, spacing=1.15)

# ============================================================ 6. MODELE PACKAGING
s = slide()
header(s, "04", "Empaqueter le modèle : un seul objet servable")
box(s, "Le bug n°1 en production, c'est un préprocessing qui diverge entre l'entraînement "
       "et l'inférence (silencieux : pas d'erreur, juste des prédictions fausses).",
    0.9, 1.7, 11.5, 0.8, 16, color=INK, spacing=1.15)
# pipeline : (libellé, remplissage, couleur du texte, bordure)
steps_pipe = [
    ("Données brutes", BLUE, NAVY, None),
    ("Feature\nengineering", WHITE, NAVY, NAVY),
    ("OneHotEncoder\n(mémorisé)", WHITE, NAVY, NAVY),
    ("RandomForest", ACCENT, WHITE, None),
    ("Probabilité", GREEN, WHITE, None),
]
x = 0.95
for i, (t, fill, txt, brd) in enumerate(steps_pipe):
    wid = 2.25 if i != 4 else 1.9
    rect(s, x, 2.75, wid, 0.95, fill, line=brd)
    box(s, t, x + 0.1, 2.85, wid - 0.2, 0.75, 13, bold=True, color=txt,
        align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0)
    if i < len(steps_pipe) - 1:
        box(s, "▸", x + wid - 0.03, 2.95, 0.35, 0.5, 16, color=ACCENT,
            align=PP_ALIGN.CENTER)
    x += wid + 0.15
box(s, "= un seul artefact  model.joblib  (Pipeline sklearn sérialisé)",
    0.95, 3.85, 11.0, 0.4, 13, bold=True, color=GREY)
bullets(s, [
    ("OneHotEncoder plutôt que get_dummies", "il mémorise les modalités au fit → prédiction fiable même pour UN employé, ou une catégorie inconnue"),
    ("Décision au seuil 0.415 (F-beta β=2), pas 0.5", "politique métier réglable sans réentraîner — priorité au rappel"),
    ("Régénéré, jamais versionné en binaire", "train.py + données + graine fixe = même modèle (reproductible)"),
], 0.95, 4.5, 11.4, size=14.5, gap=0.12)

# ============================================================ 7. API
s = slide()
header(s, "05", "API FastAPI : un contrat typé à la frontière")
bullets(s, [
    ("Validation Pydantic à l'entrée", "25 champs typés/bornés/énumérés → réponse 422 explicite AVANT le modèle"),
    ("Documentation qui ne peut pas mentir", "Swagger généré depuis le même contrat que la validation"),
    ("Endpoints clairs", "/health · POST /predict · POST /predict/batch · GET /predictions"),
    ("Chargement du modèle au démarrage", "l'artefact est désérialisé une fois (lifespan), pas à chaque appel"),
], 0.95, 1.85, 7.1, size=15, gap=0.16)
rect(s, 8.5, 1.95, 4.0, 4.35, INK)
box(s, "POST /predict", 8.75, 2.15, 3.6, 0.4, 14, bold=True, color=ACCENT, font="Courier New")
box(s, 'X-API-Key: •••••', 8.75, 2.62, 3.6, 0.3, 12, color=BLUE, font="Courier New")
box(s, '{ "age": 41, "heure_supp":\n  "Oui", ... }',
    8.75, 2.98, 3.6, 0.7, 12, color=WHITE, font="Courier New", spacing=1.1)
box(s, "▼", 10.2, 3.75, 0.6, 0.4, 16, color=ACCENT, align=PP_ALIGN.CENTER)
box(s, '{ "probabilite_depart":\n  0.70,\n  "prediction": 1,\n  "seuil": 0.415 }',
    8.75, 4.2, 3.6, 1.4, 13, color=GREEN, font="Courier New", spacing=1.15)
box(s, "Réponse tracée en base", 8.75, 5.75, 3.6, 0.4, 12, bold=True, color=BLUE)

# ============================================================ 8. POSTGRES
s = slide()
header(s, "06", "PostgreSQL : la mémoire de l'activité du modèle")
box(s, "Toute interaction passe par la base — c'est ce qui rend le modèle auditable et "
       "supervisable.",
    0.9, 1.7, 11.5, 0.6, 16, color=INK, spacing=1.1)
# deux tables
rect(s, 0.95, 2.5, 3.7, 3.4, WHITE, line=RGBColor(0xE1, 0xE5, 0xEC))
rect(s, 0.95, 2.5, 3.7, 0.55, NAVY)
box(s, "employes", 1.1, 2.6, 3.4, 0.4, 15, bold=True, color=WHITE, font="Courier New")
box(s, "id_employee (PK)\n25 champs métier\na_quitte_l_entreprise (label)",
    1.15, 3.25, 3.3, 2.4, 13, color=INK, font="Courier New", spacing=1.3)
box(s, "le dataset de référence\n(1470 employés)", 1.15, 5.15, 3.3, 0.6, 12, color=GREY)

rect(s, 4.95, 2.5, 3.9, 3.4, WHITE, line=RGBColor(0xE1, 0xE5, 0xEC))
rect(s, 4.95, 2.5, 3.9, 0.55, ACCENT)
box(s, "predictions", 5.1, 2.6, 3.6, 0.4, 15, bold=True, color=WHITE, font="Courier New")
box(s, "id (PK)\n25 champs (snapshot input)\nprobabilite · prediction\nseuil · version · created_at",
    5.15, 3.25, 3.6, 2.4, 13, color=INK, font="Courier New", spacing=1.3)
box(s, "la trace de chaque appel", 5.15, 5.4, 3.6, 0.4, 12, color=GREY)

bullets(s, [
    ("Mixin : 25 champs déclarés une fois", "alignement des tables garanti par construction"),
    ("Snapshot, pas de clé étrangère", "on peut scorer un employé absent du dataset"),
    ("À quoi ça sert", "audit d'une alerte · détection de dérive · réentraînement futur"),
], 9.15, 2.6, 3.5, size=13, gap=0.2)

# ============================================================ 9. TESTS (chart)
fig, ax = plt.subplots(figsize=(5.0, 3.1), dpi=200)
mods = ["preprocessing", "app/main", "schemas", "orm", "GLOBAL"]
cov = [100, 100, 100, 100, 94.7]
cols = ["#243A63"] * 4 + ["#E86A33"]
ax.barh(mods, cov, color=cols)
ax.set_xlim(0, 100)
ax.axvline(90, color="#2E8B57", ls="--", lw=1.5)
ax.text(90, 4.7, " plancher 90 %", color="#2E8B57", fontsize=8, va="center")
for i, v in enumerate(cov):
    ax.text(v - 2, i, f"{v:.0f}%" if v == int(v) else f"{v}%", color="white",
            ha="right", va="center", fontsize=9, fontweight="bold")
ax.invert_yaxis()
ax.set_xlabel("Couverture (%)", fontsize=9)
ax.tick_params(labelsize=9)
for sp in ["top", "right"]:
    ax.spines[sp].set_visible(False)
fig.tight_layout()
cov_png = ICI / "_cov.png"
fig.savefig(cov_png, bbox_inches="tight")
plt.close(fig)

s = slide()
header(s, "07", "Tests & couverture : la fiabilité est mesurée")
s.shapes.add_picture(str(cov_png), Inches(0.9), Inches(1.9), width=Inches(6.3))
bullets(s, [
    ("54 tests unitaires ET fonctionnels", "preprocessing · modèle · API · base · entraînement"),
    ("Couverture 94,7 %, plancher 90 % en CI", "la couverture ne peut pas régresser silencieusement"),
    ("On teste surtout ce qui doit échouer", "8 rejets 422 · base indisponible · artefact manquant"),
    ("Un bug rencontré = un test écrit", "la fragilité pandas 3 est devenue un test permanent"),
    ("Reproductibilité testée", "réentraîner redonne le même seuil (0.415)"),
], 7.6, 1.95, 5.1, size=13.5, gap=0.14)

# ============================================================ 10. SECURITE
s = slide()
header(s, "08", "Sécurité & secrets")
bullets(s, [
    ("Authentification par clé d'API", "en-tête X-API-Key exigé sur les endpoints de prédiction ; /health public"),
    ("Comparaison en temps constant", "secrets.compare_digest — pas de fuite par timing"),
    ("Aucun secret dans le dépôt", ".env.example ne contient que des placeholders"),
    ("Secrets externalisés", "DATABASE_URL & API_KEY dans Render ; RENDER_DEPLOY_HOOK côté GitHub"),
    ("Moindre privilège pour le déploiement", "GitHub ne reçoit qu'un « bouton » à presser, pas les accès Render"),
], 0.95, 1.9, 7.2, size=15, gap=0.16)
rect(s, 8.5, 1.95, 4.0, 4.3, NAVY)
box(s, "API_KEY  ≠  X-API-Key", 8.75, 2.25, 3.6, 0.5, 15, bold=True, color=ACCENT, font=HEAD)
box(s, "Côté serveur : la clé attendue (Render).", 8.75, 2.9, 3.6, 0.7, 13, color=WHITE, spacing=1.15)
box(s, "Côté client : la clé présentée à chaque appel.", 8.75, 3.55, 3.6, 0.7, 13, color=WHITE, spacing=1.15)
rect(s, 8.75, 4.35, 3.4, 0.02, NAVY2)
box(s, "« La combinaison enregistrée dans le coffre vs celle qu'on tape au clavier. »",
    8.75, 4.5, 3.6, 1.4, 13, color=BLUE, spacing=1.2)

# ============================================================ 11. MODELE RESULTATS (chart)
try:
    import sys
    sys.path.insert(0, str(ICI.parent))
    import joblib
    import pandas as pd
    from ml.preprocessing import charger_donnees, separer_X_y

    pipe = joblib.load(ICI.parent / "ml" / "model.joblib")
    X, _ = separer_X_y(charger_donnees(ICI.parent / "data" / "raw"))
    stable = X.iloc[[0]].copy()
    stable.loc[:, ["heure_supplementaires"]] = "Non"
    stable.loc[:, ["satisfaction_employee_equipe",
                   "satisfaction_employee_environnement",
                   "satisfaction_employee_nature_travail",
                   "satisfaction_employee_equilibre_pro_perso"]] = 4
    stable.loc[:, ["revenu_mensuel"]] = 12000
    p_stable = float(pipe.predict_proba(stable)[0, 1])
except Exception:
    p_stable = 0.12

fig, ax = plt.subplots(figsize=(5.0, 3.1), dpi=200)
names = ["Profil stable\n(satisfait, pas d'heures sup)", "Profil A\n(équilibré)",
         "Profil B\n(junior, insatisfait)"]
vals = [p_stable, 0.70, 0.91]
colors = ["#2E8B57" if v < 0.415 else "#E86A33" for v in vals]
ax.bar(names, vals, color=colors, width=0.6)
ax.axhline(0.415, color="#1B2A4A", ls="--", lw=1.5)
ax.text(2.4, 0.44, "seuil 0.415", color="#1B2A4A", fontsize=8)
for i, v in enumerate(vals):
    ax.text(i, v + 0.02, f"{v:.2f}", ha="center", fontsize=10, fontweight="bold")
ax.set_ylim(0, 1)
ax.set_ylabel("Probabilité de départ", fontsize=9)
ax.tick_params(labelsize=8)
for sp in ["top", "right"]:
    ax.spines[sp].set_visible(False)
fig.tight_layout()
mod_png = ICI / "_model.png"
fig.savefig(mod_png, bbox_inches="tight")
plt.close(fig)

s = slide()
header(s, "09", "Le modèle en action : un outil d'alerte")
s.shapes.add_picture(str(mod_png), Inches(0.9), Inches(1.9), width=Inches(6.3))
bullets(s, [
    ("Détecte ~7 partants sur 10 (rappel 0,71)", "priorité métier : mieux vaut une fausse alerte qu'un départ raté"),
    ("Le seuil trie les profils", "au-dessus de 0.415 → alerte RH ; le curseur est réglable"),
    ("Cause n°1 d'attrition : les heures sup.", "confirmée en amont par 3 méthodes d'importance"),
    ("Ce n'est pas un oracle", "outil d'aide ; l'humain reste dans la boucle (model card)"),
], 7.6, 2.05, 5.1, size=13.5, gap=0.18)

# ============================================================ 12. INCIDENT
s = slide(NAVY)
box(s, "Un incident en production, résolu en méthode", 0.8, 0.62, 11.5, 0.8, 28, bold=True,
    color=WHITE, font=HEAD)
box(s, "Tous les tests verts, /health OK… mais /predict renvoyait 500 en prod.",
    0.8, 1.45, 11.5, 0.5, 16, color=ACCENT)
steps = [
    ("1 · Reproduire", "Rejouer la logique en local contre la vraie base Neon → ça marche.\n→ le code, le modèle et la base sont hors de cause."),
    ("2 · Lire la traceback", "Logs Render : ModuleNotFoundError: psycopg2.\nSQLAlchemy chargeait le mauvais driver."),
    ("3 · Comprendre la cause", "DATABASE_URL = postgresql:// (sans driver) → psycopg2 par défaut,\nabsent (on utilise psycopg v3)."),
    ("4 · Corriger la classe de bug", "Immédiat : +psycopg dans l'URL.\nDurable : normaliser_url() + 2 tests → le piège ne revient jamais."),
]
y = 2.15
for t, d in steps:
    rect(s, 0.8, y, 11.7, 1.02, NAVY2)
    box(s, t, 1.0, y + 0.12, 3.0, 0.8, 15, bold=True, color=ACCENT, font=HEAD,
        anchor=MSO_ANCHOR.MIDDLE)
    box(s, d, 4.1, y + 0.1, 8.2, 0.85, 12.5, color=WHITE, spacing=1.05,
        anchor=MSO_ANCHOR.MIDDLE)
    y += 1.14
box(s, "Un code vert en test peut échouer en prod sur une seule différence d'environnement.",
    0.8, 6.85, 11.7, 0.4, 13, bold=True, color=BLUE, align=PP_ALIGN.CENTER)

# ============================================================ 13. BILAN
s = slide(NAVY)
rect(s, 0, 0, W.inches, 0.28, ACCENT)
box(s, "Bilan : du notebook au service en production", 0.8, 0.85, 11.5, 0.9, 30, bold=True,
    color=WHITE, font=HEAD)
livr = [
    ("Dépôt Git", "structuré · 7 tags · 9 PR"),
    ("API déployée", "FastAPI + Swagger, live"),
    ("Tests", "54 tests · 94,7 % couverture"),
    ("PostgreSQL", "traçabilité complète"),
    ("CI/CD", "tests + déploiement par tag"),
    ("Documentation", "README · model card · schéma"),
]
x, y = 0.8, 1.95
for i, (t, d) in enumerate(livr):
    cx = x + (i % 3) * 4.0
    cy = y + (i // 3) * 1.5
    rect(s, cx, cy, 3.75, 1.3, NAVY2)
    box(s, "✓ " + t, cx + 0.2, cy + 0.16, 3.4, 0.4, 15, bold=True, color=ACCENT, font=HEAD)
    box(s, d, cx + 0.2, cy + 0.62, 3.4, 0.6, 12.5, color=WHITE, spacing=1.05)
rect(s, 0.8, 5.2, 11.7, 0.02, NAVY2)
box(s, "Perspective : un tableau de bord de suivi de dérive alimenté par la table predictions.",
    0.8, 5.45, 11.7, 0.5, 15, color=BLUE)
box(s, "Merci — questions ?", 0.8, 6.1, 11.7, 0.7, 26, bold=True, color=WHITE, font=HEAD)

out = ICI / "Senant_Olivier_soutenance_deploiement.pptx"
prs.save(str(out))
print(f"Deck écrit : {out}  ({len(prs.slides.__iter__.__self__._sldIdLst)} slides)")
